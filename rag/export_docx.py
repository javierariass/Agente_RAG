"""Exporta una respuesta del chat a un documento Word (.docx)."""

import io
import re
from datetime import datetime

import docx
from docx.shared import Pt, RGBColor

ACCENT = RGBColor(0x00, 0x3B, 0x5C)
MUTED = RGBColor(0x6B, 0x7A, 0x82)

# Divide una linea en: citas [1][2], negritas **texto** y texto normal.
_TOKEN_RE = re.compile(r"(\[\d+\]|\*\*[^*]+\*\*)")
_BULLET_RE = re.compile(r"^\s*[-*•]\s+")
_NUMBERED_RE = re.compile(r"^\s*\d+[.)]\s+")
_HEADING_RE = re.compile(r"^\s*(#{1,4})\s+")


def _add_inline(paragraph, text):
    for token in _TOKEN_RE.split(text):
        if not token:
            continue
        if _TOKEN_RE.fullmatch(token) and token.startswith("["):
            run = paragraph.add_run(token)
            run.font.superscript = True
            run.font.bold = True
            run.font.color.rgb = ACCENT
        elif token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token)


def build_docx(question, answer, citations):
    """Devuelve los bytes del .docx.

    `answer` puede traer marcas de cita [n] y un markdown sencillo (negritas,
    listas, titulos). `citations` es una lista de dicts con `n`, `title`,
    `page` y `kind`, en el mismo orden en el que se numeraron.
    """
    document = docx.Document()
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = document.add_heading("Respuesta del Agente RAG", level=1)
    title.runs[0].font.color.rgb = ACCENT

    meta = document.add_paragraph()
    meta_run = meta.add_run(datetime.now().strftime("Generado el %d/%m/%Y a las %H:%M"))
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = MUTED

    if question:
        q = document.add_paragraph()
        label = q.add_run("Pregunta: ")
        label.bold = True
        q.add_run(question).italic = True

    for raw_line in (answer or "").splitlines():
        line = raw_line.rstrip()
        if not line.strip():
            continue
        heading = _HEADING_RE.match(line)
        if heading:
            level = min(len(heading.group(1)) + 1, 4)
            p = document.add_heading(level=level)
            _add_inline(p, line[heading.end():])
        elif _BULLET_RE.match(line):
            p = document.add_paragraph(style="List Bullet")
            _add_inline(p, _BULLET_RE.sub("", line, count=1))
        elif _NUMBERED_RE.match(line):
            p = document.add_paragraph(style="List Number")
            _add_inline(p, _NUMBERED_RE.sub("", line, count=1))
        else:
            _add_inline(document.add_paragraph(), line.strip())

    if citations:
        heading = document.add_heading("Fuentes", level=2)
        heading.runs[0].font.color.rgb = ACCENT
        for cit in citations:
            p = document.add_paragraph()
            num = p.add_run(f"[{cit.get('n')}] ")
            num.bold = True
            num.font.color.rgb = ACCENT
            p.add_run(str(cit.get("title") or cit.get("source") or ""))
            details = []
            if cit.get("page"):
                details.append(f"pág. {cit['page']}")
            if cit.get("kind") == "upload":
                details.append("archivo adjunto")
            elif cit.get("source"):
                details.append(str(cit["source"]))
            if details:
                extra = p.add_run("  —  " + " · ".join(details))
                extra.font.size = Pt(9)
                extra.font.color.rgb = MUTED

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
