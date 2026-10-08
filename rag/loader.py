import csv
import io
import os
import shutil
import subprocess
import tempfile
from concurrent.futures import ThreadPoolExecutor

import docx
import pptx
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".txt", ".doc"}

# Formatos aceptados para los archivos que el usuario adjunta en el chat. Son
# los mismos del indice mas hojas de calculo (que no se indexan en data/).
UPLOAD_EXTENSIONS = SUPPORTED_EXTENSIONS | {".xlsx", ".xlsm", ".csv"}


def iter_document_paths(base_path):
    for root, _, files in os.walk(base_path):
        for file in files:
            ext = os.path.splitext(file)[1].lower()
            if ext in SUPPORTED_EXTENSIONS:
                yield os.path.join(root, file)


def read_document(path):
    """Texto completo del documento (sin informacion de pagina)."""
    return "\n".join(text for _, text in read_document_parts(path))


def read_document_parts(path):
    """Devuelve una lista de tuplas (pagina, texto).

    Para PDF hay una tupla por pagina (numerada desde 1) para que las citas
    puedan enlazar a la pagina exacta. El resto de formatos devuelve una sola
    tupla con pagina None.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".pdf":
        return read_pdf_pages(path)
    if ext in (".xlsx", ".xlsm"):
        return [(None, read_xlsx(path))]
    if ext == ".csv":
        return [(None, read_csv(path))]
    return [(None, _read_single(path, ext))]


def _read_single(path, ext):
    if ext == ".docx":
        return read_docx(path)
    if ext == ".pptx":
        return read_pptx(path)
    if ext == ".txt":
        return read_txt(path)
    if ext == ".doc":
        return read_doc(path)
    return ""


def read_doc(path):
    """Lee un .doc antiguo (OLE2) convirtiendolo con LibreOffice/antiword si
    estan instalados. Si no hay convertidor disponible devuelve texto vacio y
    el documento simplemente se omite (no rompe el servicio)."""
    if shutil.which("antiword"):
        try:
            result = subprocess.run(
                ["antiword", path],
                capture_output=True, text=True, timeout=180,
            )
            if result.returncode == 0 and result.stdout.strip():
                return result.stdout
        except Exception:
            pass

    converter = shutil.which("soffice") or shutil.which("libreoffice")
    if not converter:
        return ""
    try:
        with tempfile.TemporaryDirectory() as tmp:
            result = subprocess.run(
                [converter, "--headless", "--convert-to", "txt:Text",
                 "--outdir", tmp, path],
                capture_output=True, text=True, timeout=180,
            )
            if result.returncode != 0:
                return ""
            txt_files = [f for f in os.listdir(tmp) if f.lower().endswith(".txt")]
            if not txt_files:
                return ""
            with open(os.path.join(tmp, txt_files[0]), "r",
                      encoding="utf-8", errors="ignore") as f:
                return f.read()
    except Exception:
        return ""

def _safe_parts(path):
    try:
        return read_document_parts(path)
    except Exception as exc:
        print(f"[rag] No se pudo leer {path}: {exc}")
        return []


def load_all_document(base_path="data"):
    """Devuelve una lista de tuplas (partes, fuente) con el contenido de cada
    documento. `partes` es la salida de `read_document_parts` y la `fuente` es
    la ruta relativa al directorio de datos."""
    paths = list(iter_document_paths(base_path))
    if not paths:
        return []

    max_workers = min(8, max(2, (os.cpu_count() or 2)))
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        contents = list(executor.map(_safe_parts, paths))

    docs = []
    for path, parts in zip(paths, contents):
        parts = [(page, text) for page, text in parts if text and text.strip()]
        if parts:
            rel = os.path.relpath(path, base_path)
            docs.append((parts, rel))
    return docs


def read_pdf(path):
    return "\n".join(text for _, text in read_pdf_pages(path))


def read_pdf_pages(path):
    reader = PdfReader(path)
    pages = []
    for number, page in enumerate(reader.pages, start=1):
        extracted = page.extract_text()
        if extracted and extracted.strip():
            pages.append((number, extracted))
    return pages


def read_docx(path):
    doc = docx.Document(path)
    return "\n".join([para.text for para in doc.paragraphs])


def read_pptx(path):
    presentation = pptx.Presentation(path)
    text = ""
    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text


def read_txt(path):
    with open(path, "r", encoding="utf-8", errors="ignore") as file:
        return file.read()


def read_xlsx(path):
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError(
            "Para leer Excel instala openpyxl (pip install openpyxl)"
        ) from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    blocks = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if v is None else str(v).strip() for v in row]
                if any(cells):
                    rows.append(" | ".join(cells).rstrip(" |"))
            if rows:
                blocks.append(f"Hoja: {sheet.title}\n\n" + "\n".join(rows))
    finally:
        workbook.close()
    return "\n\n".join(blocks)


def read_csv(path):
    with open(path, "rb") as f:
        raw = f.read()
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for row in csv.reader(io.StringIO(text), dialect):
        cells = [c.strip() for c in row]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)
